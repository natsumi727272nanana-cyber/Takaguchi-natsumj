"""Excelレポートの出力。"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .metrics import KPI_LABELS

YEN = '#,##0;[Red]-#,##0'
PCT = '0.0%;[Red]-0.0%'
INT = '#,##0'
FORMATS = {
    "sales": YEN, "avg_ticket": YEN, "labor_cost": YEN, "hpb_cost": YEN, "hpb_cost_per_new": YEN,
    "airpay_sales": YEN, "airpay_diff": YEN, "sales_prev_year": YEN, "target_gap": YEN,
    "visits": INT, "new_customers": INT, "repeat_customers": INT, "new_cohort": INT, "new_returned": INT,
    "new_return_rate": PCT, "nomination_rate": PCT, "spa_rate": PCT, "labor_cost_ratio": PCT,
    "sales_mom": PCT, "sales_yoy": PCT,
}
EXTRA_LABELS = {
    "month": "月", "store": "店舗", "staff": "スタイリスト", "labor_cost": "人件費", "hpb_cost": "HPB掲載費",
    "airpay_sales": "Airペイ決済額", "airpay_diff": "Airペイ差額", "sales_prev_year": "前年同月売上",
    "sales_mom": "売上 前月比", "sales_yoy": "売上 前年同月比", "new_cohort": "新規（再来率の分母）",
    "new_returned": "うち再来", "target_gap": "目標まで",
}
HEADER_FILL = PatternFill("solid", fgColor="DDE4EE")


def _label(col: str) -> str:
    return KPI_LABELS.get(col, EXTRA_LABELS.get(col, col))


def _prev_month(month: str, n: int) -> str:
    return (pd.Period(month, "M") - n).strftime("%Y-%m")


def _latest_matured(df: pd.DataFrame, month: str, keys: list[str]) -> pd.DataFrame:
    """対象月以前で、新規再来率が確定している最新月の値を返す。"""
    d = df[(df["month"] <= month) & df["new_return_rate"].notna()].sort_values("month")
    d = d.groupby(keys, as_index=False).last()
    return d[keys + ["month", "new_return_rate"]].rename(columns={"month": "new_return_month"})


def executive_summary(store_kpi: pd.DataFrame, month: str, config: dict) -> pd.DataFrame:
    stores = config["stores"]
    target = config["targets"]["monthly_sales"]
    window = config["rules"]["repeat_window_days"]
    matured = _latest_matured(store_kpi, month, ["store"]).set_index("store")
    columns = {}
    for sid, name in stores.items():
        s = store_kpi[store_kpi["store"] == sid].set_index("month")
        for tag, m in (("当月", month), ("前月", _prev_month(month, 1)), ("前年同月", _prev_month(month, 12))):
            col = []
            for k in config["rules"]["executive_kpis"]:
                if k == "new_return_rate" and tag == "当月":
                    # 当月の新規はまだ再来判定できないので、確定済みの最新月を出す
                    col.append(matured["new_return_rate"].get(sid))
                else:
                    col.append(s[k].get(m) if k in s and m in s.index else None)
            sales = s["sales"].get(m) if m in s.index else None
            col.append(target - sales if (tag == "当月" and sales is not None) else None)
            columns[(name, tag)] = col
    index = [_label(k) for k in config["rules"]["executive_kpis"]] + ["目標まで（売上）"]
    df = pd.DataFrame(columns, index=index)
    note = {sid: matured["new_return_month"].get(sid, "未確定") for sid in stores}
    df.attrs["notes"] = [
        f"対象月: {month}　目標: 月{target:,.0f}円（{config['rules']['sales_tax']}）",
        "新規再来率の『当月』欄は、" + "、".join(f"{stores[s]}={m}" for s, m in note.items())
        + f" に来店した新規客の{window}日以内再来率（当月新規はまだ判定できないため）",
        "人件費率が空欄の月は data/manual/monthly_costs.csv が未入力",
    ]
    return df


def _format_sheet(ws, df: pd.DataFrame, start_row: int = 1, index: bool = False):
    for cell in ws[start_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    offset = 1 if index else 0
    for j, col in enumerate(df.columns, start=1 + offset):
        fmt = FORMATS.get(col)
        if fmt:
            for row in ws.iter_rows(min_row=start_row + 1, min_col=j, max_col=j):
                row[0].number_format = fmt
        ws.column_dimensions[get_column_letter(j)].width = 14
    if index:
        ws.column_dimensions["A"].width = 22
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1 + offset)


def write_report(path: Path, month: str, config: dict, store_kpi: pd.DataFrame, staff_kpi: pd.DataFrame,
                 load_reports: list, warnings: list[str]):
    stores = config["stores"]
    path.parent.mkdir(parents=True, exist_ok=True)
    exec_df = executive_summary(store_kpi, month, config)

    staff_cols = ["store", "staff", "sales", "visits", "new_customers", "avg_ticket",
                  "nomination_rate", "spa_rate"]
    staff_now = staff_kpi[staff_kpi["month"] == month][staff_cols]
    staff_now = staff_now.merge(_latest_matured(staff_kpi, month, ["store", "staff"]), on=["store", "staff"], how="left")
    staff_now = staff_now.sort_values(["store", "staff"])
    staff_now["store"] = staff_now["store"].map(stores)

    trend_cols = ["month", "store", "sales", "sales_mom", "sales_yoy", "visits", "new_customers",
                  "repeat_customers", "avg_ticket", "new_return_rate", "nomination_rate", "spa_rate",
                  "labor_cost", "labor_cost_ratio", "hpb_cost", "hpb_cost_per_new"]
    order = {sid: i for i, sid in enumerate(stores)}
    trend = store_kpi[trend_cols].sort_values(["store", "month"], key=lambda c: c.map(order) if c.name == "store" else c)
    trend["store"] = trend["store"].map(stores)

    check = store_kpi.loc[trend.index, ["month", "store", "sales", "airpay_sales", "airpay_diff"]]
    check["store"] = check["store"].map(stores)

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        notes = exec_df.attrs["notes"]
        flat = exec_df.copy()
        flat.columns = [f"{s}\n{t}" for s, t in flat.columns]
        flat.to_excel(xw, sheet_name="幹部サマリー", startrow=len(notes) + 1)
        ws = xw.sheets["幹部サマリー"]
        for i, n in enumerate(notes, start=1):
            ws.cell(row=i, column=1, value=n)
        hdr = len(notes) + 2
        for cell in ws[hdr]:
            cell.font, cell.fill = Font(bold=True), HEADER_FILL
            cell.alignment = Alignment(wrap_text=True)
        keys = config["rules"]["executive_kpis"] + ["target_gap"]
        for i, k in enumerate(keys):
            for cell in ws[hdr + 1 + i][1:]:
                cell.number_format = FORMATS.get(k, INT)
        ws.column_dimensions["A"].width = 22
        for j in range(2, flat.shape[1] + 2):
            ws.column_dimensions[get_column_letter(j)].width = 14

        for name, df in (("スタイリスト別", staff_now), ("店舗別推移", trend), ("売上突合", check)):
            df.rename(columns={c: _label(c) for c in df.columns}).to_excel(xw, sheet_name=name, index=False)
            fmt_df = df.copy()
            _format_sheet(xw.sheets[name], fmt_df)
        ws = xw.sheets["スタイリスト別"]
        ws.cell(row=1, column=staff_now.columns.get_loc("new_return_month") + 1, value="再来率の対象月")

        rows = [{"ファイル": r.file, "行数": r.rows, "対応した列": ", ".join(f"{k}←{v}" for k, v in r.matched.items()),
                 "見つからない項目": ", ".join(r.missing), "読み捨てた列（値は未保存）": ", ".join(r.dropped_columns)}
                for r in load_reports]
        pd.DataFrame(rows).to_excel(xw, sheet_name="データチェック", index=False)
        ws = xw.sheets["データチェック"]
        for i, w in enumerate(warnings, start=len(rows) + 3):
            ws.cell(row=i, column=1, value="⚠ " + w)
        for col in "ABCDE":
            ws.column_dimensions[col].width = 40
